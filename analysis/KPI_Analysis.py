import os
import glob
import pymysql
from openpyxl import load_workbook
import re
import time
import logging
import datetime
import ast
import scipy.stats as stats
from scipy.stats import shapiro
import numpy as np


class DBInfo:
    def __init__(self, host):
        self.host = host['host']
        self.user = host['user']
        self.password = host['password']
        self.database = host['database']
        self.con = self.connect_to_database()

    def connect_to_database(self):
        con = pymysql.connect(
            host=self.host,
            user=self.user,
            password=self.password,
            database=self.database,
        )
        return con

    def execute_sql(self, sql, args=None):
        cur = self.con.cursor()
        cur.execute(sql, args)
        result = cur.fetchall()
        cur.close()
        return result

    def execute_query(self, table, columns, conditions, group_by=None, order_by=None):
        sql = f"SELECT {', '.join(columns)} FROM {table} WHERE {' AND '.join(conditions)}"
        if group_by:
            sql += f" GROUP BY {group_by}"
        if order_by:
            sql += f" ORDER BY {order_by}"
        return self.execute_sql(sql)

    def get_list(self, table, column, timeColumn, searchStartTime_str, searchEndTime_str, additional_conditions=None):
        conditions = [timeColumn + f" BETWEEN '{searchStartTime_str}' AND '{searchEndTime_str}'"]
        if additional_conditions:
            conditions.extend(additional_conditions)
        sql = f"SELECT DISTINCT {column} FROM {table} WHERE {' AND '.join(conditions)}"
        return [item[0] for item in self.execute_sql(sql)]

    def insert_data(self, table, data):
        cur = self.con.cursor()
        for row in data:
            placeholders = ', '.join(['%s'] * len(row))
            columns = ', '.join(row.keys())
            update_clause = ', '.join([f"{col}=VALUES({col})" for col in row.keys()])
            sql = f"""
                INSERT INTO {table} ({columns}) 
                VALUES ({placeholders})
                ON DUPLICATE KEY UPDATE {update_clause}
            """
            cur.execute(sql, list(row.values()))
        self.con.commit()
        cur.close()


    def close(self):
        self.con.close()


class Analyzer:
    def __init__(self, db):
        self.db = db

    def cumData_by_equip(self, searchStart, searchEnd, columns, group_by, order_by):
        conditions = [
            f"workdt BETWEEN '{searchStart}' AND '{searchEnd}'",
            "oper_old = '5600'"
        ]
        return self.db.execute_query('db_yielddetail', columns, conditions, group_by, order_by)

    def prime_by_equip(self, searchStart, searchEnd, columns, group_by, order_by):
        conditions = [
            f"date_val BETWEEN '{searchStart}' AND '{searchEnd}'",
        ]
        return self.db.execute_query('db_primeyieldet', columns, conditions, group_by, order_by)

    def retest_by_equip(self, searchStart, searchEnd, columns, group_by, order_by):
        conditions = [
            f"date_val BETWEEN '{searchStart}' AND '{searchEnd}'",
        ]
        return self.db.execute_query('db_et_retest', columns, conditions, group_by, order_by)

    def dcfa_by_equip(self, searchStart, searchEnd, columns, group_by, order_by):
        conditions = [
            f"workdt BETWEEN '{searchStart}' AND '{searchEnd}'",
        ]
        return self.db.execute_query('db_dcfa', columns, conditions, group_by, order_by)

    def ETRawData(self, workdtStart, workdtEnd, dateValStart, dateValEnd):
        # 查询ET Yield信息
        db = self.db
        ETcumYield_byEquip = self.cumData_by_equip(
            workdtStart,
            workdtEnd,
            columns=["main_equip_id as equip", "sum(out_qty)/sum(in_qty) * 100 as c_yield"],
            group_by="main_equip_id",
            order_by="main_equip_id asc"
        )
        ETqty_byEquip = self.cumData_by_equip(
            workdtStart,
            workdtEnd,
            columns=["main_equip_id as equip", "sum(in_qty) as qty"],
            group_by="main_equip_id",
            order_by="main_equip_id asc"
        )
        ETPrimeYield_byEquip = self.prime_by_equip(
            dateValStart,
            dateValEnd,
            columns=["equip_id as equip", "sum(bin1_cnt)/sum(test_cnt)*100 as p_yield"],
            group_by="equip_id",
            order_by="equip_id asc"
        )
        # 增加ET RETEST Rate@20250108
        ETRetestRate_byEquip = self.retest_by_equip(
            workdtStart,
            workdtEnd,
            columns=["equip_id as equip", "sum(retest_cnt)/sum(test_cnt)*100 as retest_rate"],
            group_by="equip",
            order_by="equip asc"
        )
        # 增加DC Fail Rate@20250108
        DCFA_byEquip = self.dcfa_by_equip(
            workdtStart,
            workdtEnd,
            columns=["equip_id as equip", "sum(dcfa_qty)/sum(in_qty)*1000000 as dcfa"],
            group_by="equip",
            order_by="equip asc"
        )
        combined_data = combine_data(ETcumYield_byEquip, ETPrimeYield_byEquip, ETqty_byEquip,ETRetestRate_byEquip, DCFA_byEquip)
        return combined_data

def combine_data(yield_data, prime_yield_data,qty_data, retest_data, dcfa_data):
    combined_data = {}
    for equip, c_yield in yield_data:
        combined_data[equip] = {
            'c_yield': c_yield,
            'p_yield': None,
            'qty': None,
            'retest': None,
            'dcfa': None
        }
    for equip, qty in qty_data:
        if equip in combined_data:
            combined_data[equip]['qty'] = qty
        else:
            combined_data[equip] = {
                'c_yield': None,
                'p_yield': None,
                'qty': qty,
                'retest': None,
                'dcfa': None
            }
    for equip, prime_yield in prime_yield_data:
        if equip in combined_data:
            combined_data[equip]['p_yield'] = prime_yield
        else:
            combined_data[equip] = {
                'c_yield': None,
                'p_yield': prime_yield,
                'qty': None,
                'retest': None,
                'dcfa': None
            }
    for equip, retest_rate in retest_data:
        if equip in combined_data:
            combined_data[equip]['retest'] = retest_rate
        else:
            combined_data[equip] = {
                'c_yield': None,
                'p_yield': None,
                'qty': None,
                'retest': retest_rate,
                'dcfa': None
            }
    for equip, dcfa in dcfa_data:
        if equip in combined_data:
            combined_data[equip]['dcfa'] = dcfa
        else:
            combined_data[equip] = {
                'c_yield': None,
                'p_yield': None,
                'qty': None,
                'retest': None,
                'dcfa': dcfa
            }
    return combined_data

def describe_data(data, workdt):
    c_yields = {device: float(item['c_yield']) for device, item in data.items()}
    p_yields = {device: float(item['p_yield']) for device, item in data.items()}
    qtys = {device: float(item['qty']) for device, item in data.items()}
    retests = {device: float(item['retest']) for device, item in data.items()}
    dcfas = {device: float(item['dcfa']) for device, item in data.items()}
    results = []
    for metric, values in zip(['c_yield', 'p_yield', 'qty', 'retest', 'dcfa'], [c_yields, p_yields, qtys, retests, dcfas]):
        values_list = list(values.values())
        # 1. 判断是否正态
        stat, p_value = shapiro(values_list)
        is_normal = p_value > 0.05
        # 2. 计算max-min值
        max_min = np.max(values_list) - np.min(values_list)
        # 3. 计算标准差
        std_dev = np.std(values_list, ddof=1)  # ddof=1 用于样本标准差
        # 4. 计算(max-min)/std
        max_min_std_ratio = max_min / std_dev
        # 5. 计算最小值及其对应设备
        min_value = np.min(values_list)
        min_device = [device for device, value in values.items() if value == min_value][0]
        # 6. 计算最大值及其对应设备
        max_value = np.max(values_list)
        max_device = [device for device, value in values.items() if value == max_value][0]
        results.append({
            'workdt': workdt,
            'data_type': metric,
            'normal': int(is_normal),
            'tolerence': round(max_min, 2),
            'stdv': round(std_dev, 2),
            'sigma': round(max_min_std_ratio, 2),
            'min_value': round(min_value, 2),
            'min_name': min_device,
            'max_value': round(max_value, 2),
            'max_name': max_device
        })
    return results


def main():
    log_path = r'C:\Users\Tengjun Zhao\Desktop\test'
    host = {
        'host': 'localhost',
        'user': 'remoteuser',
        'password': 'password',
        'database': 'cmsalpha',
    }

    try:
        # 获取昨天以及三天前的日期，格式是YYYYMMDD/YYYY-MM-DD
        workdt = (datetime.datetime.now()).strftime('%Y%m%d')
        workdtEnd = (datetime.datetime.now() - datetime.timedelta(days=1)).strftime('%Y%m%d')
        workdtStart = (datetime.datetime.now() - datetime.timedelta(days=3)).strftime('%Y%m%d')
        dateValEnd = (datetime.datetime.now() - datetime.timedelta(days=1)).strftime('%Y-%m-%d')
        dateValStart = (datetime.datetime.now() - datetime.timedelta(days=3)).strftime('%Y-%m-%d')
        mydb = DBInfo(host)
        myAnalyzer = Analyzer(mydb)
        etData = myAnalyzer.ETRawData(workdtStart, workdtEnd, dateValStart, dateValEnd)
        etData.pop('2MTV01', None)
        # 对数据进行分析
        data_form = describe_data(etData, workdt)
        mydb.insert_data('db_kpi_describe', data_form)
        print(data_form)
    except Exception as e:
        logging.error("Error occurred", exc_info=True)
    finally:
        mydb.close()

if __name__ == '__main__':
    main()
